from __future__ import annotations

from math import cos, pi, sin
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("09-project-presentation.pptx")

NAVY = "0B1020"
PANEL = "151C31"
PANEL_2 = "1D2742"
WHITE = "F5F7FF"
MUTED = "A8B2CB"
CYAN = "39D8FF"
PURPLE = "8B7CFF"
GREEN = "5DE2A5"
AMBER = "FFC857"
RED = "FF6B7A"
GRID = "2B3656"
INK = "111827"

FONT = "Aptos"
FONT_DISPLAY = "Aptos Display"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, radius=0.12, line=None, line_width=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(line_width)
    if radius:
        shape.adjustments[0] = min(0.22, radius)
    return shape


def circle(slide, x, y, d, fill, line=None, line_width=1):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(line_width)
    return shape


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=20,
    color=WHITE,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    fit=True,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    if fit:
        frame.fit_text(font_family=font, max_size=Pt(size))
    p = frame.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)
    return box


def rich_text(slide, runs, x, y, w, h, size=20, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(0)
    frame.margin_top = frame.margin_bottom = Pt(0)
    p = frame.paragraphs[0]
    p.alignment = align
    for value, color, bold in runs:
        r = p.add_run()
        r.text = value
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    return box


def line(slide, x1, y1, x2, y2, color=GRID, width=1.5, dash=False, arrow=False):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    if dash:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        connector.line.end_arrowhead = True
    return connector


def pill(slide, label, x, y, w, fill=PANEL_2, color=CYAN):
    rect(slide, x, y, w, 0.34, fill, radius=0.18)
    text(
        slide,
        label.upper(),
        x,
        y + 0.01,
        w,
        0.27,
        10,
        color,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def number_badge(slide, value, x, y, fill=CYAN, color=NAVY):
    circle(slide, x, y, 0.42, fill)
    text(
        slide,
        str(value),
        x,
        y + 0.005,
        0.42,
        0.34,
        14,
        color,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def base_slide(prs, title, section, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(NAVY)
    pill(slide, section, 0.66, 0.42, 1.55)
    text(slide, title, 0.66, 0.89, 11.7, 0.7, 27, WHITE, True, FONT_DISPLAY)
    line(slide, 0.66, 1.62, 12.67, 1.62, GRID, 1)
    text(
        slide,
        f"{number:02d}",
        12.18,
        7.02,
        0.48,
        0.22,
        10,
        MUTED,
        True,
        align=PP_ALIGN.RIGHT,
    )
    return slide


def node(slide, label, subtitle, x, y, w, h, accent=CYAN, fill=PANEL):
    rect(slide, x, y, w, h, fill, radius=0.08, line=GRID)
    rect(slide, x, y, 0.06, h, accent, radius=0)
    text(slide, label, x + 0.25, y + 0.18, w - 0.45, 0.32, 17, WHITE, True)
    text(slide, subtitle, x + 0.25, y + 0.58, w - 0.45, h - 0.72, 11, MUTED)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(NAVY)

    for i in range(7):
        circle(
            slide,
            9.0 + cos(i * pi / 3.5) * (0.75 + i * 0.12),
            2.75 + sin(i * pi / 3.5) * (0.75 + i * 0.12),
            0.16 + (i % 3) * 0.05,
            CYAN if i % 2 == 0 else PURPLE,
        )
    for i in range(3):
        line(slide, 9.78, 3.53, 9.15 + i * 0.72, 2.15 + i * 0.9, GRID, 1.2)
    circle(slide, 9.22, 2.95, 1.18, PANEL_2, line=CYAN, line_width=2)
    text(
        slide,
        "100M",
        9.22,
        3.27,
        1.18,
        0.3,
        20,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    text(
        slide,
        "голосов",
        9.22,
        3.61,
        1.18,
        0.2,
        10,
        MUTED,
        align=PP_ALIGN.CENTER,
    )

    pill(slide, "Архитектура решения", 0.74, 0.67, 2.2)
    text(
        slide,
        "Минутный\nТВ-опрос",
        0.74,
        1.34,
        7.4,
        1.75,
        44,
        WHITE,
        True,
        FONT_DISPLAY,
    )
    text(
        slide,
        "Как принять массовый пик голосов, не потерять результат\nи сохранить приватность зрителей",
        0.78,
        3.38,
        7.2,
        0.96,
        20,
        MUTED,
    )
    rect(slide, 0.78, 5.42, 6.42, 0.86, PANEL, radius=0.08, line=GRID)
    rich_text(
        slide,
        [
            ("Главная идея: ", WHITE, True),
            ("Redis принимает пик, PostgreSQL сохраняет правду.", CYAN, False),
        ],
        1.08,
        5.69,
        5.85,
        0.34,
        16,
    )
    text(slide, "FULL-STACK MVP · FASTAPI · REDIS · POSTGRESQL · REACT", 0.78, 6.8, 8.8, 0.2, 10, MUTED, True)


def add_problem_slide(prs):
    slide = base_slide(prs, "Задача — принять массовый пик голосов", "Контекст", 2)
    text(slide, "100 млн", 0.72, 2.04, 3.05, 0.62, 38, WHITE, True, FONT_DISPLAY)
    text(slide, "голосов за 60 секунд", 0.76, 2.69, 3.2, 0.36, 17, MUTED)
    text(slide, "≈ 1,7 млн RPS", 0.76, 3.34, 3.25, 0.5, 27, CYAN, True)
    text(slide, "ориентир для архитектуры,\nне результат локального теста", 0.76, 3.92, 3.3, 0.63, 13, MUTED)

    line(slide, 4.38, 1.98, 4.38, 6.36, GRID, 1)
    items = [
        ("01", "Быстро", "Не обращаться к SQL на каждом голосе", CYAN),
        ("02", "Надёжно", "Не считать повтор и восстанавливать итог", GREEN),
        ("03", "Приватно", "Без регистрации и хранения сырого IP", AMBER),
    ]
    y = 2.14
    for num, heading, body, accent in items:
        text(slide, num, 4.77, y + 0.05, 0.5, 0.28, 12, accent, True)
        text(slide, heading, 5.38, y, 2.05, 0.34, 18, WHITE, True)
        text(slide, body, 7.35, y, 4.72, 0.6, 13, MUTED)
        if y < 4.9:
            line(slide, 4.77, y + 0.98, 12.24, y + 0.98, GRID, 1)
        y += 1.24
    rect(slide, 4.76, 6.18, 7.48, 0.54, PANEL_2, radius=0.06)
    text(slide, "Решение: быстрый приём отдельно от надёжного хранения.", 4.97, 6.33, 6.98, 0.22, 13, WHITE, True)


def add_architecture_slide(prs):
    slide = base_slide(prs, "Два контура: горячий путь и источник правды", "Архитектура", 3)
    node(slide, "Зрители", "QR → браузер", 0.72, 2.32, 1.65, 1.0, CYAN)
    node(slide, "Nginx", "TLS · proxy", 2.76, 2.32, 1.54, 1.0, PURPLE)
    node(slide, "FastAPI", "stateless API", 4.72, 2.32, 1.85, 1.0, CYAN)
    node(slide, "Redis", "dedup · counters · cache", 7.22, 2.03, 2.24, 1.56, GREEN)
    node(slide, "PostgreSQL", "journal · CRUD · rebuild", 10.05, 2.03, 2.38, 1.56, PURPLE)

    line(slide, 2.37, 2.82, 2.76, 2.82, CYAN, 2, arrow=True)
    line(slide, 4.30, 2.82, 4.72, 2.82, CYAN, 2, arrow=True)
    line(slide, 6.57, 2.56, 7.22, 2.56, GREEN, 2.4, arrow=True)
    line(slide, 9.46, 2.82, 10.05, 2.82, PURPLE, 1.8, dash=True, arrow=True)
    text(slide, "HOT PATH", 6.61, 2.16, 0.54, 0.19, 9, GREEN, True)
    text(slide, "async journal", 9.48, 2.5, 0.55, 0.3, 9, MUTED, True, align=PP_ALIGN.CENTER)

    node(slide, "Администратор", "CRUD · QR · live result", 0.72, 4.83, 2.28, 1.06, AMBER)
    line(slide, 3.0, 5.35, 4.72, 3.15, AMBER, 1.5, arrow=True)

    rect(slide, 4.72, 4.62, 7.71, 1.54, PANEL, radius=0.08, line=GRID)
    text(slide, "Почему это масштабируется", 5.03, 4.86, 2.8, 0.34, 16, WHITE, True)
    facts = [
        ("API", "без локального состояния"),
        ("Redis", "поглощает короткий пик"),
        ("Postgres", "восстанавливает итог"),
    ]
    for i, (head, body) in enumerate(facts):
        x = 5.05 + i * 2.4
        text(slide, head, x, 5.36, 0.86, 0.25, 12, [CYAN, GREEN, PURPLE][i], True)
        text(slide, body, x, 5.66, 2.1, 0.36, 11, MUTED)
    text(slide, "Редкие admin-операции и массовое голосование не конкурируют за один путь.", 0.74, 6.64, 10.8, 0.26, 13, WHITE, True)


def add_vote_flow_slide(prs):
    slide = base_slide(prs, "Как принимается один голос", "Hot path", 4)
    steps = [
        ("1", "Запрос", "question · option · vid"),
        ("2", "Проверка", "статус · время · вариант"),
        ("3", "Dedup", "хеш question + vid"),
        ("4", "Lua", "проверка + инкремент"),
        ("5", "Ответ", "201 или 409"),
    ]
    x_positions = [0.95, 3.34, 5.73, 8.12, 10.51]
    for i, ((num, head, body), x) in enumerate(zip(steps, x_positions)):
        number_badge(slide, num, x + 0.68, 2.08, CYAN if i < 3 else GREEN)
        if i < len(steps) - 1:
            line(slide, x + 1.11, 2.29, x_positions[i + 1] + 0.67, 2.29, GRID, 2, arrow=True)
        text(slide, head, x, 2.72, 1.78, 0.28, 15, WHITE, True, align=PP_ALIGN.CENTER)
        text(slide, body, x - 0.12, 3.09, 2.02, 0.64, 11, MUTED, align=PP_ALIGN.CENTER)

    rect(slide, 1.02, 4.19, 11.28, 1.54, PANEL, radius=0.08, line=GRID)
    text(slide, "АТОМАРНАЯ ОПЕРАЦИЯ", 1.36, 4.5, 1.7, 0.2, 10, GREEN, True)
    text(slide, "EXISTS dedup?", 3.03, 4.42, 1.52, 0.34, 18, WHITE, True)
    line(slide, 4.56, 4.62, 5.18, 4.62, GRID, 2, arrow=True)
    rect(slide, 5.2, 4.35, 3.2, 0.58, PANEL_2, radius=0.08, line=GREEN)
    text(slide, "SET dedup EX ttl", 5.2, 4.52, 3.2, 0.22, 15, GREEN, True, align=PP_ALIGN.CENTER)
    line(slide, 8.42, 4.62, 9.03, 4.62, GRID, 2, arrow=True)
    text(slide, "HINCRBY", 9.08, 4.42, 1.42, 0.34, 18, WHITE, True)
    text(slide, "повтор → 409", 3.02, 5.12, 1.54, 0.22, 11, RED, True)
    text(slide, "новый голос → 201", 9.08, 5.12, 1.8, 0.22, 11, GREEN, True)
    text(slide, "Lua атомарно отмечает зрителя и увеличивает счётчик.", 1.03, 6.22, 8.6, 0.3, 16, WHITE, True)
    pill(slide, "O(1) на голос", 10.62, 6.14, 1.68, PANEL_2, CYAN)


def add_data_slide(prs):
    slide = base_slide(prs, "Где хранятся данные и результаты", "Данные", 6)
    text(slide, "REDIS", 0.74, 1.96, 2.0, 0.3, 14, GREEN, True)
    text(slide, "POSTGRESQL", 7.14, 1.96, 2.0, 0.3, 14, PURPLE, True)

    redis_items = [
        ("question:{id}", "публичная карточка"),
        ("vote:{q}:{hash}", "dedup с TTL"),
        ("results:{q}:{shard}", "счётчики вариантов"),
    ]
    pg_items = [
        ("question", "статус и окно эфира"),
        ("question_option", "варианты ответа"),
        ("vote", "append-only журнал"),
        ("question_result", "финальный снимок"),
    ]
    for i, (key, body) in enumerate(redis_items):
        y = 2.42 + i * 1.03
        rect(slide, 0.74, y, 5.38, 0.78, PANEL, radius=0.06, line=GRID)
        text(slide, key, 1.02, y + 0.16, 2.25, 0.25, 14, WHITE, True)
        text(slide, body, 3.42, y + 0.18, 2.35, 0.23, 12, MUTED)
    for i, (key, body) in enumerate(pg_items):
        y = 2.42 + i * 0.77
        circle(slide, 7.15, y + 0.08, 0.31, PURPLE)
        text(slide, key, 7.65, y + 0.01, 2.08, 0.27, 14, WHITE, True)
        text(slide, body, 9.73, y + 0.02, 2.48, 0.27, 12, MUTED)
        if i < 3:
            line(slide, 7.31, y + 0.39, 7.31, y + 0.84, GRID, 1)

    line(slide, 6.63, 2.28, 6.63, 5.58, GRID, 1)
    rect(slide, 0.74, 5.94, 11.48, 0.63, PANEL_2, radius=0.08)
    rich_text(
        slide,
        [
            ("Приватность: ", WHITE, True),
            ("cookie vid — браузер, не человек · IP хранится только как salted hash · ", MUTED, False),
            ("ограничение: cookie ≠ антибот.", AMBER, True),
        ],
        1.03,
        6.14,
        10.95,
        0.28,
        13,
    )


def add_demo_slide(prs):
    slide = base_slide(prs, "Демонстрация проекта", "Демо", 5)
    rect(slide, 0.72, 1.96, 8.15, 4.78, PANEL, radius=0.08, line=CYAN, line_width=1.5)
    circle(slide, 4.12, 3.42, 1.18, CYAN)
    text(
        slide,
        "▶",
        4.19,
        3.67,
        1.0,
        0.42,
        28,
        NAVY,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    text(
        slide,
        "ВСТАВЬТЕ ВИДЕО ПРОЕКТА",
        1.74,
        4.92,
        6.12,
        0.32,
        16,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )
    text(
        slide,
        "Рекомендуемая длительность: 60–90 секунд",
        1.74,
        5.39,
        6.12,
        0.28,
        12,
        MUTED,
        align=PP_ALIGN.CENTER,
    )

    text(slide, "Что показать", 9.39, 2.02, 2.4, 0.36, 18, WHITE, True)
    points = [
        ("01", "Админ создаёт опрос"),
        ("02", "Зритель голосует по QR"),
        ("03", "Результат обновляется live"),
        ("04", "Повторный голос отклонён"),
    ]
    for i, (num, body) in enumerate(points):
        y = 2.73 + i * 0.82
        number_badge(slide, num, 9.39, y, CYAN if i < 2 else GREEN)
        text(slide, body, 10.02, y + 0.06, 2.55, 0.42, 13, WHITE)
    rect(slide, 9.37, 6.23, 3.2, 0.51, PANEL_2, radius=0.06)
    text(slide, "Видео заменяет длинный рассказ о UI", 9.59, 6.37, 2.76, 0.2, 11, MUTED, True, align=PP_ALIGN.CENTER)


def add_quality_slide(prs):
    slide = base_slide(prs, "Как проверяется решение", "Качество", 7)
    checks = [
        ("Unit", "логика backend"),
        ("Lint + Build", "Python и TypeScript"),
        ("HTTP", "API-сценарии"),
        ("E2E", "путь пользователя"),
    ]
    for i, (head, body) in enumerate(checks):
        x = 0.98 + i * 3.03
        circle(slide, x + 0.86, 2.05, 0.58, GREEN)
        text(slide, "✓", x + 0.86, 2.16, 0.58, 0.3, 20, NAVY, True, align=PP_ALIGN.CENTER)
        if i < 3:
            line(slide, x + 1.45, 2.34, x + 3.0, 2.34, GRID, 2)
        text(slide, head, x, 2.86, 2.3, 0.3, 15, WHITE, True, align=PP_ALIGN.CENTER)
        text(slide, body, x, 3.23, 2.3, 0.3, 11, MUTED, align=PP_ALIGN.CENTER)

    rect(slide, 0.76, 4.08, 11.84, 1.62, PANEL, radius=0.08, line=GRID)
    text(slide, "Критические сценарии", 1.04, 4.36, 2.45, 0.31, 16, WHITE, True)
    scenarios = [
        "параллельный повтор",
        "невалидный вариант без записи",
        "два браузера за одним IP",
        "rebuild результата",
    ]
    for i, item in enumerate(scenarios):
        x = 1.03 + (i % 2) * 5.74
        y = 4.91 + (i // 2) * 0.46
        circle(slide, x, y + 0.03, 0.16, CYAN)
        text(slide, item, x + 0.28, y, 5.05, 0.26, 12, MUTED)
    rect(slide, 4.92, 6.16, 3.5, 0.5, PANEL_2, radius=0.06)
    text(slide, "make verify → одна команда", 4.92, 6.3, 3.5, 0.21, 13, CYAN, True, align=PP_ALIGN.CENTER)


def add_risks_slide(prs):
    slide = base_slide(prs, "Что готово и что остаётся", "Итог", 8)
    risks = [
        ("01", "Нагрузка", "масштаб нужно подтвердить тестами", AMBER),
        ("02", "Dedup", "cookie защищает от повтора, не от бота", CYAN),
        ("03", "Надёжность", "async-журнал нужно вынести в durable queue", RED),
        ("04", "Безопасность", "для production нужны OIDC и RBAC", PURPLE),
    ]
    for i, (num, head, body, accent) in enumerate(risks):
        col = i % 2
        row = i // 2
        x = 0.74 + col * 6.12
        y = 2.02 + row * 1.24
        rect(slide, x, y, 5.64, 0.98, PANEL, radius=0.06, line=GRID)
        text(slide, num, x + 0.24, y + 0.2, 0.4, 0.2, 10, accent, True)
        text(slide, head, x + 0.81, y + 0.16, 1.78, 0.27, 15, WHITE, True)
        text(slide, body, x + 2.57, y + 0.16, 2.76, 0.55, 11, MUTED)

    rect(slide, 0.74, 4.79, 5.64, 1.27, PANEL_2, radius=0.06, line=GREEN)
    text(slide, "ГОТОВО", 1.03, 5.05, 0.92, 0.2, 10, GREEN, True)
    text(slide, "Full-stack MVP, Docker,\nавтотесты и live-результаты", 2.0, 4.99, 3.92, 0.6, 15, WHITE, True)
    rect(slide, 6.86, 4.79, 5.64, 1.27, PANEL_2, radius=0.06, line=CYAN)
    text(slide, "ДАЛЕЕ", 7.15, 5.05, 0.92, 0.2, 10, CYAN, True)
    text(slide, "Нагрузочное доказательство,\nнаблюдаемость и hardening", 8.12, 4.99, 3.92, 0.6, 15, WHITE, True)
    text(
        slide,
        "MVP демонстрирует правильные архитектурные границы.",
        0.76,
        6.49,
        11.74,
        0.32,
        17,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Минутный ТВ-опрос"
    prs.core_properties.subject = "Архитектура высоконагруженного голосования"
    prs.core_properties.author = "Crypto Enthusiast"
    prs.core_properties.keywords = "FastAPI, Redis, PostgreSQL, architecture"
    add_title_slide(prs)
    add_problem_slide(prs)
    add_architecture_slide(prs)
    add_vote_flow_slide(prs)
    add_demo_slide(prs)
    add_data_slide(prs)
    add_quality_slide(prs)
    add_risks_slide(prs)
    prs.save(OUT)
    print(f"Saved {len(prs.slides)} slides to {OUT}")


if __name__ == "__main__":
    build()
